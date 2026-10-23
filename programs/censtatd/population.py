"""Census and Statistics Department — population."""


def population() -> list:
    """Mid-year Hong Kong population by sex, annual series.

    Provides:
        - mid-year population by sex (male / female / both), annual from 1961

    Coverage:
        - Geography: HK-wide
        - Time:      historical series, annual

    Datasets:   hk-censtatd-tablechart-110-02001
    Department: Census and Statistics Department
    Endpoint:   https://www.censtatd.gov.hk/api/get.php?id=110-02001&lang=en&full_series=1
    Format:     JSON

    Usage:
        >>> rows = population()               # list[dict]
        >>> [r for r in rows if r["sv"] == "POP_TOTAL"]

    Questions:
        - "how many people live in Hong Kong?"
        - "Hong Kong population by year and sex"

    Quirks:
        - needs a browser-like User-Agent (bare urllib gets 403)
        - record fields: sv/svDesc (variable), period (year), figure
    """
    import json
    import urllib.request

    url = ("https://www.censtatd.gov.hk/api/get.php"
           "?id=110-02001&lang=en&full_series=1")
    req = urllib.request.Request(
        url, headers={"User-Agent":
                      "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.loads(resp.read().decode("utf-8"))["dataSet"]


def language_use(sheet: str = "A107e") -> dict:
    """2021 Population Census language-use tables (usual language / ability).

    Provides:
        - workbook metadata (source URL, sheet)
        - rows: language / dialect with 2011 / 2016 / 2021 counts and %
          (shape depends on sheet; A107e = usual spoken language;
           A111e = ability to speak as usual / another / total %)

    Coverage:
        - Geography: HK-wide (population aged 5+)
        - Time:      census years 2011 / 2016 / 2021 (static workbook)

    Datasets:   hk-censtatd-census21c-d5212101
    Department: Census and Statistics Department
    Endpoint:   https://www.censtatd.gov.hk/en/data/stat_report/product/D5212101/att/D5212101E2021XXXXE.xlsx
    Format:     XLSX

    Usage:
        >>> usual = language_use("A107e")
        >>> able = language_use("A111e")
        >>> usual["rows"][0]

    Questions:
        - "How many Hong Kong people speak English / Putonghua?"
        - "Hong Kong usual spoken language census"
        - "香港人慣用語言 廣州話 英語 普通話 比例"
        - "code-switching / 中英夾雜 statistics" (no official measure — say so)

    Quirks:
        - no official dataset measures code-switching/code-mixing (中英夾雜);
          give census/THS proxies and state the gap
        - hk.sh test can UnicodeDecodeError on these XLSX URLs — curl + openpyxl
        - THS Report No.76 (C0000086) has situation/frequency tables; this
          function reads the census main tables workbook only
        - census excludes mute persons; sheet titles sit above the header block
    """
    import io

    import openpyxl
    import requests

    url = ("https://www.censtatd.gov.hk/en/data/stat_report/product/"
           "D5212101/att/D5212101E2021XXXXE.xlsx")
    resp = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=90)
    resp.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True,
                                read_only=True)
    name = sheet if sheet in wb.sheetnames else sheet
    if name not in wb.sheetnames:
        raise ValueError(
            f"sheet {sheet!r} not in workbook; try A107e / A111e / A120e")
    ws = wb[name]

    if name == "A107e":
        rows = []
        for i, raw in enumerate(ws.iter_rows(values_only=True)):
            if i < 6:
                continue
            lang = raw[0]
            if lang is None or str(lang).startswith("Note") or str(
                    lang).startswith("("):
                if lang and str(lang).startswith("Note"):
                    break
                continue
            rows.append({
                "language": lang,
                "persons_2011": raw[1],
                "pct_2011": raw[2],
                "persons_2016": raw[3],
                "pct_2016": raw[4],
                "persons_2021": raw[5],
                "pct_2021": raw[6],
            })
            if lang == "Total":
                break
        return {
            "source_url": url,
            "sheet": name,
            "title": "Usual spoken language, population aged 5+",
            "years": [2011, 2016, 2021],
            "rows": rows,
            "note": ("Code-switching / 中英夾雜 is NOT measured in official "
                     "census or THS datasets."),
        }

    if name == "A111e":
        rows = []
        for i, raw in enumerate(ws.iter_rows(values_only=True)):
            if i < 6:
                continue
            lang = raw[0]
            if lang is None:
                continue
            s = str(lang)
            if s.startswith("Note") or s.startswith("("):
                break
            rows.append({
                "language": lang,
                "usual_pct_2011": raw[1],
                "usual_pct_2016": raw[2],
                "usual_pct_2021": raw[3],
                "another_pct_2011": raw[4],
                "another_pct_2016": raw[5],
                "another_pct_2021": raw[6],
                "total_able_pct_2011": raw[7],
                "total_able_pct_2016": raw[8],
                "total_able_pct_2021": raw[9],
            })
        return {
            "source_url": url,
            "sheet": name,
            "title": ("Ability to speak selected languages/dialects "
                      "(usual / another / total %), aged 5+"),
            "years": [2011, 2016, 2021],
            "rows": rows,
            "note": ("Code-switching / 中英夾雜 is NOT measured in official "
                     "census or THS datasets."),
        }

    # Generic fallback: first non-empty cells per row
    generic = []
    for raw in ws.iter_rows(values_only=True):
        vals = [c for c in raw if c is not None]
        if vals:
            generic.append(vals)
    return {
        "source_url": url,
        "sheet": name,
        "title": f"Raw rows from sheet {name}",
        "rows": generic[:200],
        "note": ("Code-switching / 中英夾雜 is NOT measured in official "
                 "census or THS datasets."),
    }
